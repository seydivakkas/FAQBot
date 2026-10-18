"""CSV/XLSX SSS aktarımı: dosya sınırı, önizleme ve atomik yeni kayıt ekleme."""
import csv
import io
import json
import zipfile
from pathlib import Path
from datetime import datetime, timezone
from openpyxl import load_workbook
from pydantic import ValidationError
from app.database import _connect
from app.matching import normalize

MAX_FILE_BYTES = 2 * 1024 * 1024
MAX_EXPANDED_BYTES = 16 * 1024 * 1024
MAX_ROWS = 500
ALIASES = {
    "question": "question", "soru": "question",
    "answer": "answer", "cevap": "answer", "yanıt": "answer",
    "category": "category", "kategori": "category",
    "alternative_questions": "alternative_questions",
    "alternatif_sorular": "alternative_questions",
    "alternatif sorular": "alternative_questions",
    "source_title": "source_title", "kaynak_basligi": "source_title",
    "source_url": "source_url", "kaynak_url": "source_url",
    "valid_until": "valid_until", "gecerlilik_tarihi": "valid_until",
    "verification_status": "verification_status", "dogrulama_durumu": "verification_status",
}


class ImportFormatError(ValueError):
    pass


class ImportConflict(ValueError):
    def __init__(self, details):
        self.details = details
        super().__init__("Dosyayı yeniden önizleyin: satırlar değişmiş veya çakışıyor.")


def _delimiter(text):
    header = text.splitlines()[0] if text.splitlines() else ""
    try:
        return csv.Sniffer().sniff(header, delimiters=",;\t").delimiter
    except csv.Error:
        return ";" if header.count(";") > header.count(",") else ","


def _parse_rows(iterator):
    try:
        first = next(iterator)
    except StopIteration:
        raise ImportFormatError("Dosya boş; başlık satırı bulunamadı.")
    columns = [ALIASES.get(str(v or "").strip().casefold().replace("\ufeff", "")) for v in first]
    defined = [v for v in columns if v]
    if len(defined) != len(set(defined)):
        raise ImportFormatError("Aynı sütun birden fazla tanımlanmış.")
    missing = {"question", "answer", "category"} - set(defined)
    if missing:
        raise ImportFormatError("Eksik sütun: " + ", ".join(sorted(missing)))
    result = []
    for number, values in enumerate(iterator, 2):
        if number > MAX_ROWS + 1:
            raise ImportFormatError(f"En fazla {MAX_ROWS} veri satırı yüklenebilir.")
        if not any(v is not None and str(v).strip() for v in values):
            continue
        record = {}
        for index, name in enumerate(columns):
            if name:
                value = values[index] if index < len(values) else ""
                record[name] = "" if value is None else str(value).strip()
        alternatives = record.get("alternative_questions", "")
        record["alternative_questions"] = [
            phrase.strip() for phrase in alternatives.split("|") if phrase.strip()
        ]
        if not record.get("valid_until"):
            record["valid_until"] = None
        if not record.get("verification_status"):
            record["verification_status"] = "unverified"
        result.append({"row_number": number, "values": record})
    return result


def parse_file(filename, content):
    ext = Path(filename or "").suffix.lower()
    if ext not in (".csv", ".xlsx"):
        raise ImportFormatError("Yalnızca CSV veya XLSX dosyası yükleyebilirsiniz.")
    if len(content) > MAX_FILE_BYTES:
        raise ImportFormatError("Dosya 2 MB sınırını aşıyor.")
    if not content:
        raise ImportFormatError("Dosya boş.")
    if ext == ".csv":
        try:
            string = content.decode("utf-8-sig")
        except UnicodeError as exc:
            raise ImportFormatError("CSV UTF-8 olarak kaydedilmelidir.") from exc
        try:
            return _parse_rows(iter(csv.reader(io.StringIO(string, newline=""),
                                               delimiter=_delimiter(string))))
        except csv.Error as exc:
            raise ImportFormatError("CSV okunamadı.") from exc
    try:
        with zipfile.ZipFile(io.BytesIO(content)) as archive:
            entries = archive.infolist()
            if len(entries) > 250 or sum(e.file_size for e in entries) > MAX_EXPANDED_BYTES:
                raise ImportFormatError("Excel arşivi açılmış boyut sınırını aşıyor.")
        book = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        try:
            if not book.worksheets:
                raise ImportFormatError("Excel dosyasında sayfa yok.")
            return _parse_rows(iter(book.worksheets[0].iter_rows(values_only=True)))
        finally:
            book.close()
    except (zipfile.BadZipFile, ValueError, OSError, KeyError) as exc:
        if isinstance(exc, ImportFormatError):
            raise
        raise ImportFormatError("Geçerli bir XLSX dosyası değil.") from exc


def analyze(records, existing, faq_model):
    """Normalize edilmiş aynı soru/alternatif ifadeleri yakalar.

    Semantik tekrar tespiti ayrı bir geliştirme aşamasıdır.
    """
    occupied = {}
    for faq in existing:
        for phrase in [faq["question"], *faq["alternative_questions"]]:
            occupied.setdefault(normalize(phrase), faq["question"])
    results = []
    for item in records:
        raw = item["values"]
        result = {
            "row_number": item["row_number"],
            "question": str(raw.get("question", ""))[:250],
            "category": str(raw.get("category", ""))[:30],
        }
        try:
            model = faq_model.model_validate(raw)
            payload = model.model_dump()
            if payload.get("valid_until") is not None:
                payload["valid_until"] = payload["valid_until"].isoformat()
            phrases = [payload["question"], *payload["alternative_questions"]]
            keys = [normalize(phrase) for phrase in phrases]
            if not all(keys):
                raise ValueError("Sorular harf veya rakam içermelidir.")
            duplicates = [occupied[key] for key in keys if key in occupied]
            if duplicates:
                result.update(status="duplicate", reason="Çakışan ifade: " + duplicates[0])
            else:
                result.update(status="ready", reason="", data=payload)
                occupied.update({key: payload["question"] for key in keys})
        except (ValidationError, ValueError) as exc:
            reason = (f"{'.'.join(map(str, exc.errors()[0]['loc']))}: "
                      f"{exc.errors()[0]['msg']}") if isinstance(exc, ValidationError) else str(exc)
            result.update(status="invalid", reason=reason)
        results.append(result)
    return results


def summarize(rows):
    return {
        "total": len(rows),
        "ready": sum(r["status"] == "ready" for r in rows),
        "duplicate": sum(r["status"] == "duplicate" for r in rows),
        "invalid": sum(r["status"] == "invalid" for r in rows),
        "rows": rows,
    }


def commit_records(records, faq_model):
    """Tamamı geçerli ve benzersizse SQLite transaction'ında INSERT yapar."""
    if not 1 <= len(records) <= MAX_ROWS:
        raise ValueError(f"1 ile {MAX_ROWS} arasında kayıt gönderilmelidir.")
    numbered = [{"row_number": i+1, "values": value} for i, value in enumerate(records)]
    with _connect() as conn:
        conn.execute("BEGIN IMMEDIATE")
        current = conn.execute("SELECT question, alternative_questions FROM faqs").fetchall()
        existing = [
            {"question": row["question"],
             "alternative_questions": json.loads(row["alternative_questions"])}
            for row in current
        ]
        checked = analyze(numbered, existing, faq_model)
        if any(r["status"] != "ready" for r in checked):
            raise ImportConflict(summarize(checked))
        for item in checked:
            row = item["data"]
            conn.execute(
                "INSERT INTO faqs (question, answer, category, alternative_questions, "
                "source_title, source_url, valid_until, verification_status, verified_at) "
                "VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?)",
                (
                    row["question"], row["answer"], row["category"],
                    json.dumps(row["alternative_questions"], ensure_ascii=False),
                    row.get("source_title", ""), row.get("source_url", ""),
                    row.get("valid_until"), row.get("verification_status", "unverified"),
                    datetime.now(timezone.utc).isoformat()
                    if row.get("verification_status") == "verified" else None,
                ),
            )
    return len(checked)
